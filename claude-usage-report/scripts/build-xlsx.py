#!/usr/bin/env python3
"""Build a styled XLSX from the enriched sessions JSON produced by run.js."""
import argparse
import json
import sys
from collections import OrderedDict
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.stderr.write("openpyxl missing. Run: pip3 install --user openpyxl\n")
    sys.exit(1)


def proj_label(cwd: str, home: str) -> str:
    """Convert an absolute project path into a short readable label.

    Default rules collapse `~/` and `~/Code/` and tag external-volume paths.
    Set CLAUDE_REPORT_STRIP_PREFIX to strip an extra leading directory under
    `~/Code/` (e.g. an org-specific monorepo subdir like "MyOrg") so labels
    stay short — without committing that name to source.
    """
    import os
    if not cwd:
        return "?"
    extra = os.environ.get("CLAUDE_REPORT_STRIP_PREFIX", "").strip("/")
    rules = []
    if extra:
        rules.extend([
            (f"/Volumes/Macintosh HD{home}/Code/{extra}/", "[ExtVol] "),
            (f"{home}/Code/{extra}/", ""),
        ])
    rules.extend([
        (f"/Volumes/Macintosh HD{home}/Code/", "[ExtVol] "),
        (f"{home}/Code/", ""),
        (f"{home}/", "~/"),
    ])
    for src, dst in rules:
        if cwd.startswith(src):
            return dst + cwd[len(src):]
    if cwd == home:
        return "~/"
    return cwd


def fmt_duration(first_ts: str, last_ts: str) -> str:
    from datetime import datetime
    a = datetime.fromisoformat(first_ts.replace("Z", "+00:00"))
    b = datetime.fromisoformat(last_ts.replace("Z", "+00:00"))
    minutes = (b - a).total_seconds() / 60
    if minutes < 1:
        return "<1m"
    if minutes < 60:
        return f"{round(minutes)}m"
    hours = minutes / 60
    if hours > 12:
        return "12h+"
    return f"{hours:.1f}h"


def fmt_tokens(n: int) -> str:
    return f"{n/1e6:.1f}M" if n >= 1e6 else f"{round(n/1e3)}K"


def weekday(date_iso: str) -> str:
    from datetime import datetime
    d = datetime.fromisoformat(date_iso)
    return ["周一", "周二", "周三", "周四", "周五", "周六", "周日"][d.weekday()]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--in", dest="inp", required=True)
    ap.add_argument("--out", required=True)
    args = ap.parse_args()

    data = json.loads(Path(args.inp).read_text(encoding="utf-8"))
    sessions = data["sessions"]
    total_cost_usd = data.get("total_cost", 0.0)
    total_tokens = data.get("total_tokens", sum(s["tokens"] for s in sessions))
    since = data.get("since", "")
    until = data.get("until", "")

    home = str(Path.home())

    # Per-session scaled cost
    def scale_cost(tokens):
        if total_tokens <= 0:
            return 0.0
        return tokens / total_tokens * total_cost_usd

    # Sort by date desc then firstTs asc (within-day chronological)
    def sort_key(s):
        d = s.get("date") or s["firstTs"][:10]
        return (d, s["firstTs"])
    sessions.sort(key=sort_key)
    sessions.reverse()  # newest day first; but within a day we want chronological
    # Re-sort: group by date desc, but inside each day chronological
    days = OrderedDict()
    for s in sessions:
        d = s.get("date") or s["firstTs"][:10]
        days.setdefault(d, []).append(s)
    # ensure within-day chronological (firstTs ascending)
    for d in days:
        days[d].sort(key=lambda s: s["firstTs"])

    # ---- Build workbook ----
    wb = Workbook()
    ws = wb.active
    ws.title = "Claude Code 使用报表"

    # Title row — count unique sessions (a multi-day session is still ONE session,
    # even though it occupies multiple rows after the cross-day split).
    unique_sessions = len({s.get("sessionFile") for s in sessions if s.get("sessionFile")})
    if unique_sessions == 0:
        unique_sessions = len(sessions)
    title = f"Claude Code 任务报表 · {since} ~ {until} · {unique_sessions} sessions · {round(total_tokens/1e6)}M tokens · ${round(total_cost_usd)}"
    ws.append([title])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=8)
    ws.cell(1, 1).font = Font(bold=True, size=14)
    ws.cell(1, 1).alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 24

    # Headers
    headers = ["Date", "星期", "Project", "Branch", "Duration", "Tokens", "Cost", "Summary"]
    ws.append(headers)

    # Styles
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4472C4")
    day_fill = PatternFill("solid", fgColor="FFF2CC")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_wrap = Alignment(horizontal="left", vertical="top", wrap_text=True)
    thin = Side(border_style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for c in ws[2]:
        c.font = header_font
        c.fill = header_fill
        c.alignment = center
        c.border = border

    # Day banners + session rows
    for date in days:
        sess = days[date]
        day_tokens = sum(s["tokens"] for s in sess)
        day_cost = sum(scale_cost(s["tokens"]) for s in sess)
        banner = f"{date}（{weekday(date)}）  ·  {len(sess)} session  ·  {day_tokens/1e6:.1f}M tokens  ·  ${day_cost:.2f}"
        ws.append([banner])
        br_row = ws.max_row
        ws.merge_cells(start_row=br_row, start_column=1, end_row=br_row, end_column=8)
        cell = ws.cell(br_row, 1)
        cell.font = Font(bold=True, size=11)
        cell.fill = day_fill
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = border

        for s in sess:
            ws.append([
                date,
                weekday(date),
                proj_label(s.get("cwd"), home),
                s.get("branch") or "-",
                fmt_duration(s["firstTs"], s["lastTs"]),
                fmt_tokens(s["tokens"]),
                f"${scale_cost(s['tokens']):.2f}",
                s.get("summary", "-"),
            ])
            for c in ws[ws.max_row]:
                c.alignment = left_wrap
                c.border = border

    # Column widths
    widths = [9, 6, 28, 22, 11, 9, 8, 60]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A3"

    out_path = Path(args.out).expanduser()
    wb.save(out_path)
    print(f"  Saved: {out_path} ({out_path.stat().st_size/1024:.1f} KB)")


if __name__ == "__main__":
    main()
